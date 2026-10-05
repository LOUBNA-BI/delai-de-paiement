import io
import math
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib.pyplot as plt
import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Délai de paiement – DGI", page_icon="💳", layout="wide")

# -----------------------------------------------------------------------------
# Paramètres juridiques / financiers par défaut
# -----------------------------------------------------------------------------
DEFAULT_LEGAL_DAYS = 60
MAX_CONTRACTUAL_DAYS = 120
ADDITIONAL_MONTH_RATE = 0.0085

# Taux directeur BAM utiles pour les données 2025 du fichier fourni.
# 2,50 % jusqu'au 19/03/2025, puis 2,25 % à compter du 20/03/2025.
BAM_RATES = pd.DataFrame(
    {
        "date_effet": pd.to_datetime(["2025-01-01", "2025-03-20"]),
        "taux": [0.025, 0.0225],
    }
)

# Jours fériés 2025 uniquement informatifs : le calcul DGI principal est en jours calendaires.
HOLIDAYS_2025 = pd.to_datetime(
    [
        "2025-01-01", "2025-01-11", "2025-01-14", "2025-03-31", "2025-04-01",
        "2025-05-01", "2025-06-06", "2025-06-07", "2025-06-27", "2025-07-30",
        "2025-08-14", "2025-08-20", "2025-08-21", "2025-10-31", "2025-11-06",
        "2025-11-18",
    ]
)


def money(x):
    return f"{x:,.2f}".replace(",", " ")


def normalize_columns(df):
    """Normalise les colonnes de la feuille 4411 du fichier source."""
    df = df.copy()
    expected = ["Date", "Journal", "Partenaire", "Factures", "Libellé", "Débit", "Crédit", "Lettrage"]
    if len(df.columns) >= 8:
        df = df.iloc[:, :8]
        df.columns = expected
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    for c in ["Débit", "Crédit"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    for c in ["Journal", "Partenaire", "Factures", "Libellé", "Lettrage"]:
        df[c] = df[c].fillna("").astype(str).str.strip()
    return df


def load_source(uploaded_file):
    """Charge la première feuille comptable. XLSM/XLSX acceptés."""
    suffix = Path(uploaded_file.name).suffix.lower()
    engine = "openpyxl" if suffix in [".xlsx", ".xlsm", ".xltx", ".xltm"] else None
    if engine is None:
        raise ValueError("Le fichier doit être .xlsx ou .xlsm.")
    df = pd.read_excel(uploaded_file, sheet_name="4411", engine=engine)
    return normalize_columns(df)


def build_allocations(raw, max_rows=None):
    """Rapproche factures et règlements par lettrage, FIFO pour les règlements partiels."""
    invoices = raw.loc[
        (raw["Journal"].str.lower() == "factures fournisseurs") & raw["Crédit"].notna()
    ].copy()
    payments = raw.loc[raw["Débit"].notna()].copy()

    inv_by_let = {}
    pay_by_let = {}
    for let, g in invoices.groupby("Lettrage", dropna=False):
        inv_by_let[str(let)] = g.sort_values("Date").to_dict("records")
    for let, g in payments.groupby("Lettrage", dropna=False):
        pay_by_let[str(let)] = g.sort_values("Date").to_dict("records")

    rows = []
    for let in sorted(inv_by_let):
        invs = inv_by_let[let]
        pays = pay_by_let.get(let, [])
        remaining = [float(x["Crédit"]) for x in invs]
        for p in pays:
            left = float(p["Débit"])
            for i, inv in enumerate(invs):
                if left <= 0.005:
                    break
                allocation = min(remaining[i], left)
                if allocation <= 0.005:
                    continue
                rows.append(
                    {
                        "Date facture": inv["Date"],
                        "Facture": inv["Factures"],
                        "Fournisseur": inv["Partenaire"],
                        "Montant facture": float(inv["Crédit"]),
                        "Lettrage": let,
                        "Mode de paiement": p["Libellé"],
                        "Banque / journal": p["Journal"],
                        "Date paiement": p["Date"],
                        "Montant payé": allocation,
                    }
                )
                remaining[i] -= allocation
                left -= allocation

    result = pd.DataFrame(rows)
    if result.empty:
        return result
    result = result.sort_values(["Date facture", "Facture", "Date paiement"]).reset_index(drop=True)
    if max_rows:
        result = result.head(max_rows).copy()
    return result


def bam_rate_at(dt):
    if pd.isna(dt):
        return float(BAM_RATES.iloc[-1]["taux"])
    d = pd.Timestamp(dt)
    applicable = BAM_RATES[BAM_RATES["date_effet"] <= d]
    if applicable.empty:
        return float(BAM_RATES.iloc[0]["taux"])
    return float(applicable.iloc[-1]["taux"])


def add_months(d, months):
    d = pd.Timestamp(d)
    return d + pd.DateOffset(months=months)


def late_months(due, payment):
    """Nombre de mois ou fraction de mois de retard, conformément au principe de fraction."""
    due = pd.Timestamp(due)
    payment = pd.Timestamp(payment)
    if payment <= due:
        return 0
    months = (payment.year - due.year) * 12 + payment.month - due.month
    anniversary = due + pd.DateOffset(months=months)
    if payment > anniversary:
        months += 1
    return max(1, months)


def first_month_end_rate(due):
    due = pd.Timestamp(due)
    first_month_end = due + pd.offsets.MonthEnd(1)
    return bam_rate_at(first_month_end)


def calculate_dgi(allocations, contractual_days=60):
    df = allocations.copy()
    if df.empty:
        return df

    contractual_days = int(contractual_days)
    contractual_days = max(0, contractual_days)
    effective_days = min(contractual_days, MAX_CONTRACTUAL_DAYS)
    df["Délai convenu (j)"] = effective_days
    df["Délai légal défaut (j)"] = DEFAULT_LEGAL_DAYS
    # À défaut de contrat : 60 jours. Avec contrat : délai convenu, plafonné à 120 jours.
    df["Délai applicable (j)"] = effective_days if contractual_days != DEFAULT_LEGAL_DAYS else DEFAULT_LEGAL_DAYS
    # L'utilisateur peut saisir 90/120 jours ; pour 60, on applique 60.
    df["Échéance"] = df["Date facture"] + pd.to_timedelta(df["Délai applicable (j)"], unit="D")
    df["Délai réel (j)"] = (df["Date paiement"] - df["Date facture"]).dt.days
    df["Retard (j)"] = (df["Date paiement"] - df["Échéance"]).dt.days.clip(lower=0)
    df["Mois / fraction retard"] = [late_months(d, p) for d, p in zip(df["Échéance"], df["Date paiement"])]
    df["Taux BAM 1er mois"] = [first_month_end_rate(d) if r > 0 else 0 for d, r in zip(df["Échéance"], df["Retard (j)"])]
    df["Taux mois supplémentaires"] = [ADDITIONAL_MONTH_RATE if m > 1 else 0 for m in df["Mois / fraction retard"]]
    df["Base amende"] = df["Montant payé"].where(df["Retard (j)"] > 0, 0.0)
    df["Amende théorique"] = df.apply(
        lambda r: r["Base amende"] * (
            r["Taux BAM 1er mois"] + max(0, r["Mois / fraction retard"] - 1) * ADDITIONAL_MONTH_RATE
        ),
        axis=1,
    )
    df["Statut DGI"] = df["Retard (j)"].apply(lambda x: "RETARD" if x > 0 else "DANS LE DÉLAI")
    df["Dépassement plafond 120 j"] = df["Délai convenu (j)"] > MAX_CONTRACTUAL_DAYS
    df["Trimestre paiement"] = "T" + df["Date paiement"].dt.quarter.astype(str) + " " + df["Date paiement"].dt.year.astype(str)
    return df


def unpaid_invoices(raw, contractual_days=60):
    invoices = raw.loc[
        (raw["Journal"].str.lower() == "factures fournisseurs") & raw["Crédit"].notna()
    ].copy()
    if invoices.empty:
        return invoices
    paid = raw.loc[raw["Débit"].notna()].groupby("Lettrage")["Débit"].sum()
    invoices["Total payé par lettrage"] = invoices["Lettrage"].map(paid).fillna(0)
    invoices["Solde estimé"] = (invoices["Crédit"] - invoices["Total payé par lettrage"]).clip(lower=0)
    invoices = invoices[invoices["Solde estimé"] > 0.005].copy()
    invoices["Échéance"] = invoices["Date"] + pd.to_timedelta(contractual_days, unit="D")
    return invoices


def applicability(ca_ht, year):
    if ca_ht <= 2_000_000:
        return "Hors champ du dispositif de sanctions (seuil CA HT ≤ 2 MDH)"
    if year == 2025:
        if ca_ht <= 10_000_000:
            return "2025 : déclaration annuelle (2 à 10 MDH CA HT)"
        if ca_ht <= 50_000_000:
            return "2025 : déclaration annuelle (10 à 50 MDH CA HT)"
        return "2025 : déclaration trimestrielle (> 50 MDH CA HT)"
    if year >= 2026:
        return "Déclaration trimestrielle (> 2 MDH CA HT)"
    if ca_ht <= 10_000_000:
        return "Phase progressive – vérifier l'exercice concerné"
    return "Phase progressive – vérifier l'exercice concerné"


def make_excel(source_file, raw, detail, ca_ht, contractual_days):
    """Ajoute une page Réponses et reconstruit la 2e page Délai de paiement."""
    source_file.seek(0)
    keep_vba = Path(source_file.name).suffix.lower() == ".xlsm"
    wb = openpyxl.load_workbook(source_file, keep_vba=keep_vba)

    if "Réponses" in wb.sheetnames:
        del wb["Réponses"]
    if "Delai de paiement" in wb.sheetnames:
        del wb["Delai de paiement"]

    ws_answer = wb.create_sheet("Réponses", 0)
    ws_dp = wb.create_sheet("Delai de paiement", 1)

    title_fill = PatternFill("solid", fgColor="1F4E78")
    head_fill = PatternFill("solid", fgColor="D9EAF7")
    late_fill = PatternFill("solid", fgColor="F4CCCC")
    thin = Side(style="thin", color="B7B7B7")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Réponses / dashboard
    ws_answer.merge_cells("A1:F1")
    ws_answer["A1"] = "RÉPONSES – ANALYSE DU DÉLAI DE PAIEMENT"
    ws_answer["A1"].fill = title_fill
    ws_answer["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws_answer["A1"].alignment = Alignment(horizontal="center")

    kpis = [
        ("Nombre de factures rapprochées", int(detail["Facture"].nunique()) if not detail.empty else 0),
        ("Montant des factures analysées", float(detail["Montant facture"].sum()) if not detail.empty else 0),
        ("Montant payé analysé", float(detail["Montant payé"].sum()) if not detail.empty else 0),
        ("Paiements en retard", int((detail["Retard (j)"] > 0).sum()) if not detail.empty else 0),
        ("Montant payé hors délai", float(detail.loc[detail["Retard (j)"] > 0, "Base amende"].sum()) if not detail.empty else 0),
        ("Amende théorique", float(detail["Amende théorique"].sum()) if not detail.empty else 0),
    ]
    for i, (label, value) in enumerate(kpis, 3):
        ws_answer.cell(i, 1, label).font = Font(bold=True)
        ws_answer.cell(i, 1).fill = head_fill
        ws_answer.cell(i, 1).border = bd
        ws_answer.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        ws_answer.cell(i, 5, value).border = bd
        if isinstance(value, float):
            ws_answer.cell(i, 5).number_format = '#,##0.00'

    ws_answer["A11"] = "Réponse DGI"
    ws_answer["A11"].font = Font(bold=True, size=11, color="1F4E78")
    responses = [
        ("Délai appliqué", f"{contractual_days} jours" if contractual_days != 60 else "60 jours à défaut d'accord"),
        ("Plafond contractuel", "120 jours maximum"),
        ("Régime 2025 selon CA HT", applicability(ca_ht, 2025)),
        ("Conclusion", "Des lignes sont en retard" if (not detail.empty and (detail["Retard (j)"] > 0).any()) else "Aucun retard sur les paiements analysés"),
    ]
    for i, (q, a) in enumerate(responses, 12):
        ws_answer.cell(i, 1, q).font = Font(bold=True)
        ws_answer.cell(i, 2, a)
        ws_answer.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws_answer.cell(i, 1).border = ws_answer.cell(i, 2).border = bd

    for c, w in enumerate([34, 25, 16, 16, 18, 18], 1):
        ws_answer.column_dimensions[get_column_letter(c)].width = w
    ws_answer.sheet_view.showGridLines = False

    # Détail
    headers = [
        "Date facture", "Facture", "Fournisseur", "Montant facture", "Lettrage",
        "Mode de paiement", "Banque / journal", "Date paiement", "Montant payé",
        "Délai convenu (j)", "Délai applicable (j)", "Échéance", "Délai réel (j)",
        "Retard (j)", "Mois / fraction retard", "Taux BAM 1er mois", "Taux mois supplémentaires",
        "Base amende", "Amende théorique", "Statut DGI", "Trimestre", "Observation"
    ]
    for j, h in enumerate(headers, 1):
        c = ws_dp.cell(1, j, h)
        c.fill = head_fill
        c.font = Font(bold=True)
        c.border = bd
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r, (_, row) in enumerate(detail.iterrows(), 2):
        vals = [
            row["Date facture"], row["Facture"], row["Fournisseur"], row["Montant facture"], row["Lettrage"],
            row["Mode de paiement"], row["Banque / journal"], row["Date paiement"], row["Montant payé"],
            row["Délai convenu (j)"], row["Délai applicable (j)"], row["Échéance"], row["Délai réel (j)"],
            row["Retard (j)"], row["Mois / fraction retard"], row["Taux BAM 1er mois"], row["Taux mois supplémentaires"],
            row["Base amende"], row["Amende théorique"], row["Statut DGI"], row["Trimestre paiement"],
            "Calculé selon loi 69-21 / DGI"
        ]
        for j, val in enumerate(vals, 1):
            c = ws_dp.cell(r, j, val)
            c.border = bd
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for j in [1, 8, 12]:
            ws_dp.cell(r, j).number_format = "dd/mm/yyyy"
        for j in [4, 9, 18, 19]:
            ws_dp.cell(r, j).number_format = '#,##0.00'
        for j in [16, 17]:
            ws_dp.cell(r, j).number_format = '0.00%'
        if row["Retard (j)"] > 0:
            for j in range(1, len(headers) + 1):
                ws_dp.cell(r, j).fill = late_fill

    if not detail.empty:
        tab = Table(displayName="Table_DP_DGI", ref=f"A1:V{len(detail)+1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws_dp.add_table(tab)

    widths = [13, 22, 25, 17, 11, 18, 20, 13, 16, 15, 17, 13, 13, 11, 16, 16, 18, 15, 15, 20, 14, 35]
    for c, w in enumerate(widths, 1):
        ws_dp.column_dimensions[get_column_letter(c)].width = w
    ws_dp.freeze_panes = "A2"
    ws_dp.sheet_view.showGridLines = False
    ws_dp.auto_filter.ref = f"A1:V{max(2, len(detail)+1)}"
    ws_dp.page_setup.orientation = "landscape"
    ws_dp.page_setup.fitToWidth = 1
    ws_dp.sheet_properties.pageSetUpPr.fitToPage = True

    # Paramètres
    if "Paramètres DGI" in wb.sheetnames:
        del wb["Paramètres DGI"]
    wp = wb.create_sheet("Paramètres DGI")
    wp.append(["Paramètre", "Valeur"])
    wp.append(["Délai légal par défaut (jours)", DEFAULT_LEGAL_DAYS])
    wp.append(["Délai contractuel maximum (jours)", MAX_CONTRACTUAL_DAYS])
    wp.append(["Taux supplémentaire par mois/fraction", ADDITIONAL_MONTH_RATE])
    wp.append([])
    wp.append(["Date effet BAM", "Taux directeur"])
    for _, rr in BAM_RATES.iterrows():
        wp.append([rr["date_effet"].to_pydatetime(), rr["taux"]])
    for row in wp.iter_rows():
        for c in row:
            c.border = bd
    wp["B4"].number_format = "0.00%"
    for r in range(7, 7 + len(BAM_RATES)):
        wp.cell(r, 1).number_format = "dd/mm/yyyy"
        wp.cell(r, 2).number_format = "0.00%"
    wp.column_dimensions["A"].width = 38
    wp.column_dimensions["B"].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# -----------------------------------------------------------------------------
# Interface
# -----------------------------------------------------------------------------
st.title("💳 Délai de paiement – Analyse DGI")
st.caption("Application Streamlit pour analyser un grand livre 4411, rapprocher factures/règlements et produire la 2e page DGI.")

with st.sidebar:
    st.header("Paramètres")
    ca_ht = st.number_input("CA HT du débiteur (DH)", min_value=0.0, value=20_000_000.0, step=500_000.0)
    contractual_days = st.number_input(
        "Délai contractuel (jours)", min_value=60, max_value=120, value=60, step=1,
        help="60 jours si aucun accord différent n'est documenté. Maximum légal : 120 jours."
    )
    max_rows = st.number_input("Nombre de lignes à afficher/exporter", min_value=1, max_value=10000, value=30, step=10)
    st.divider()
    st.markdown("**Référence : loi n°69-21 / DGI**")
    st.markdown("60 j par défaut • 120 j maximum • BAM + 0,85 %/mois supplémentaire")

uploaded = st.file_uploader("Charger le fichier Excel 4411 (.xlsm ou .xlsx)", type=["xlsm", "xlsx"])

if not uploaded:
    st.info("Chargez `Compte 4411-DP.xlsm` pour commencer.")
    st.stop()

try:
    raw = load_source(uploaded)
except Exception as e:
    st.error(f"Impossible de lire la feuille 4411 : {e}")
    st.stop()

full_detail = build_allocations(raw)
detail = full_detail.head(int(max_rows)).copy()
if not detail.empty:
    detail = calculate_dgi(detail, int(contractual_days))

# -----------------------------------------------------------------------------
# Réponses à la première page
# -----------------------------------------------------------------------------
st.subheader("1. Réponses / diagnostic")
if detail.empty:
    st.warning("Aucune facture fournisseur rapprochée avec un règlement n'a été trouvée.")
else:
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Factures", f"{detail['Facture'].nunique():,}")
    c2.metric("Montant payé", f"{money(detail['Montant payé'].sum())} DH")
    c3.metric("Retards", int((detail["Retard (j)"] > 0).sum()))
    c4.metric("Jours max de retard", int(detail["Retard (j)"].max()))
    c5.metric("Amende théorique", f"{money(detail['Amende théorique'].sum())} DH")

    conclusion = "⚠️ Des paiements dépassent le délai applicable." if (detail["Retard (j)"] > 0).any() else "✅ Les paiements analysés sont dans le délai applicable."
    st.info(conclusion)
    st.write(f"**Régime DGI 2025 selon le CA HT saisi :** {applicability(ca_ht, 2025)}")
    st.caption("Le calcul porte sur les lignes rapprochées. Les factures non lettrées/non payées sont présentées séparément ci-dessous.")

# -----------------------------------------------------------------------------
# 2e page
# -----------------------------------------------------------------------------
st.subheader("2. Deuxième page – détail DGI")
if detail.empty:
    st.warning("Pas de détail à afficher.")
else:
    show_cols = [
        "Date facture", "Facture", "Fournisseur", "Montant facture", "Date paiement", "Montant payé",
        "Délai applicable (j)", "Échéance", "Délai réel (j)", "Retard (j)", "Mois / fraction retard",
        "Taux BAM 1er mois", "Base amende", "Amende théorique", "Statut DGI"
    ]
    st.dataframe(detail[show_cols], use_container_width=True, hide_index=True)

# -----------------------------------------------------------------------------
# Factures non payées / non rapprochées
# -----------------------------------------------------------------------------
st.subheader("3. Factures non soldées détectées")
unpaid = unpaid_invoices(raw, int(contractual_days))
if unpaid.empty:
    st.success("Aucune facture non soldée détectée sur les lettrages disponibles.")
else:
    st.dataframe(
        unpaid[["Date", "Factures", "Partenaire", "Crédit", "Lettrage", "Total payé par lettrage", "Solde estimé", "Échéance"]],
        use_container_width=True,
        hide_index=True,
    )

# -----------------------------------------------------------------------------
# Graphique matplotlib
# -----------------------------------------------------------------------------
st.subheader("4. Visualisation")
if not detail.empty:
    chart = detail.groupby("Fournisseur", as_index=False)["Retard (j)"].max().sort_values("Retard (j)", ascending=False).head(15)
    fig, ax = plt.subplots(figsize=(10, 4.5))
    ax.bar(chart["Fournisseur"], chart["Retard (j)"])
    ax.set_ylabel("Retard maximum (jours)")
    ax.set_title("Retard maximum par fournisseur")
    ax.tick_params(axis="x", rotation=45)
    fig.tight_layout()
    st.pyplot(fig, clear_figure=True)

# -----------------------------------------------------------------------------
# Export Excel
# -----------------------------------------------------------------------------
st.subheader("5. Export Excel")
if st.button("Préparer le fichier Excel complet", type="primary"):
    try:
        # Recalcule le détail selon le nombre choisi, et conserve le fichier source complet.
        export_detail = full_detail.head(int(max_rows)).copy()
        export_detail = calculate_dgi(export_detail, int(contractual_days))
        # UploadedFile doit être repositionné avant openpyxl.
        uploaded.seek(0)
        xlsx_bytes = make_excel(uploaded, raw, export_detail, ca_ht, int(contractual_days))
        st.download_button(
            "⬇️ Télécharger le fichier Excel",
            data=xlsx_bytes,
            file_name="Compte_4411_DP_DGI_analyse.xlsm" if uploaded.name.lower().endswith(".xlsm") else "Compte_4411_DP_DGI_analyse.xlsx",
            mime="application/vnd.ms-excel.sheet.macroEnabled.12" if uploaded.name.lower().endswith(".xlsm") else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as e:
        st.error(f"Erreur pendant la génération Excel : {e}")

st.divider()
st.caption("Outil d'analyse : les résultats doivent être rapprochés des pièces justificatives, du CA HT de l'assujetti, des conditions contractuelles et des données DGI avant dépôt.")
